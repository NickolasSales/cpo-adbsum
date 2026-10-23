"""
Importacao da relacao de ministros, a partir de XLSX ou CSV.

Construido sobre a planilha oficial ("Relatorio Ministros 2026.xlsx"), sem
exigir que ninguem a edite antes de enviar. O que o arquivo real tem, e que
um leitor ingenuo nao aguentaria:

    titulo da relacao nas linhas 2 a 4 (celulas mescladas)
    cabecalho so na linha 6, com uma linha em branco logo abaixo
    uma coluna sem cabecalho entre "Nome" e "Inscricao Conamad"
    "Data Nascimento" mesclada sobre TRES colunas: dia, mes e ano
    "IDADE" e formula (=2026-ano) — derivada, nao e importada
    "nº" e a posicao na relacao (1..N), nao matricula
    titulo abreviado e com caixa variada: EV, Ev, PR, PRª, Prª, Mis

Por isso o cabecalho e PROCURADO nas primeiras linhas, e nao assumido na
primeira; e a data de nascimento aceita uma celula de data, um texto
dd/mm/aaaa ou tres celulas numericas debaixo de um cabecalho so.

Duas fases, como na importacao de alunos:

    analisar()   le e classifica. NAO escreve nada.
    confirmar()  refaz a analise sobre as linhas cruas guardadas na sessao e
                 grava, numa unica transacao, somente as linhas NOVAS.

Ministro ja cadastrado nunca e alterado pela importacao. Se o arquivo traz
dados diferentes, o preview mostra "dados divergentes" e quais campos; a
correcao e feita pela tela de edicao, de proposito.
"""

import csv
import datetime
import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from xml.etree.ElementTree import ParseError

from django.db import IntegrityError, transaction
from django.utils import timezone

from audit.models import AuditEvent
from audit.services import record
from common.exceptions import DomainError
from ministers import services
from ministers.models import Minister, MinisterialTitle, MinisterSex, MinisterSource
from ministers.normalizacao import (
    TAMANHO_MAXIMO_REGISTRO,
    chave_de_nome,
    chave_de_texto,
    normalizar_nome,
    normalizar_registro,
    registro_valido,
)

logger = logging.getLogger("cpo.ministers")

TAMANHO_MAXIMO_BYTES = 5 * 1024 * 1024
MAX_LINHAS = 3000
EXTENSOES_ACEITAS = (".csv", ".xlsx")

# Quantas linhas do topo sao varridas atras do cabecalho. A relacao oficial
# o tem na linha 6; vinte cobre qualquer titulo razoavel acima dele.
LINHAS_PROCURANDO_CABECALHO = 20

# Cabecalho normalizado (chave_de_texto) -> campo interno.
#
# "matricula" sozinho NAO esta aqui: na guia em PDF existem duas matriculas
# (campo e CONAMAD), e adivinhar qual delas a coluna quer dizer seria
# inventar semantica. Uma coluna assim aparece como ignorada no preview.
ALIAS_COLUNAS = {
    "n": "sequencia",
    "no": "sequencia",
    "n o": "sequencia",
    "numero": "sequencia",
    "nome": "nome",
    "nome completo": "nome",
    "nome do ministro": "nome",
    "ministro": "nome",
    "inscricao conamad": "conamad",
    "inscricao da conamad": "conamad",
    "matricula conamad": "conamad",
    "registro conamad": "conamad",
    "conamad": "conamad",
    "matricula campo": "campo",
    "matricula do campo": "campo",
    "matricula de campo": "campo",
    "titulo": "titulo",
    "titulo ministerial": "titulo",
    "cargo": "cargo",
    "funcao": "cargo",
    "consagracao": "consagracao",
    "ano consagracao": "consagracao",
    "ano de consagracao": "consagracao",
    "ano da consagracao": "consagracao",
    "data nascimento": "nascimento",
    "data de nascimento": "nascimento",
    "nascimento": "nascimento",
    "sexo": "sexo",
    "congregacao": "congregacao",
    "observacoes": "observacoes",
    "observacao": "observacoes",
}

# Colunas reconhecidas e deliberadamente nao importadas.
COLUNAS_DERIVADAS = {"idade": "calculada a partir do ano de nascimento"}

ROTULO_DO_CAMPO = {
    "sequencia": "Nº na relação",
    "nome": "Nome",
    "conamad": "Inscrição CONAMAD",
    "campo": "Matrícula do campo",
    "titulo": "Título ministerial",
    "cargo": "Cargo / função",
    "consagracao": "Ano de consagração",
    "nascimento": "Data de nascimento",
    "sexo": "Sexo",
    "congregacao": "Congregação",
    "observacoes": "Observações",
}

TITULOS = {
    "ev": MinisterialTitle.EVANGELISTA,
    "evang": MinisterialTitle.EVANGELISTA,
    "evangelista": MinisterialTitle.EVANGELISTA,
    "pr": MinisterialTitle.PASTOR,
    "pastor": MinisterialTitle.PASTOR,
    # "PRª" e "Prª": o ordinal feminino vira "a" na normalizacao.
    "pra": MinisterialTitle.PASTORA,
    "pastora": MinisterialTitle.PASTORA,
    "mis": MinisterialTitle.MISSIONARIO,
    "miss": MinisterialTitle.MISSIONARIO,
    "missionario": MinisterialTitle.MISSIONARIO,
    "missionaria": MinisterialTitle.MISSIONARIO,
    "missionario a": MinisterialTitle.MISSIONARIO,
}

# A coluna Sexo da relacao mistura dois esquemas: H/M (homem/mulher) e
# M/F (masculino/feminino). "H" e "F" sao inequivocos. "M" aparece tanto em
# ministros quanto em ministras — inclusive em uma Pastora e em todas as
# missionarias —, entao nao da para saber o que ele quer dizer em cada linha.
# Ele entra em branco, com aviso, para ser corrigido na edicao.
SEXOS = {
    "h": MinisterSex.MASCULINO,
    "homem": MinisterSex.MASCULINO,
    "masculino": MinisterSex.MASCULINO,
    "masc": MinisterSex.MASCULINO,
    "f": MinisterSex.FEMININO,
    "feminino": MinisterSex.FEMININO,
    "fem": MinisterSex.FEMININO,
    "mulher": MinisterSex.FEMININO,
}
SEXOS_AMBIGUOS = {"m"}


class RowStatus:
    NOVO = "NOVO"
    EXISTENTE = "JÁ EXISTE"
    DIVERGENTE = "JÁ EXISTE — DADOS DIVERGENTES"
    DUPLICADO = "DUPLICADO NA PLANILHA"
    OBRIGATORIO = "CAMPO OBRIGATÓRIO AUSENTE"
    CONFLITO = "MATRÍCULA CONFLITANTE"
    INVALIDA = "LINHA INVÁLIDA"


STATUS_EXISTENTES = frozenset({RowStatus.EXISTENTE, RowStatus.DIVERGENTE})
STATUS_RECUSADOS = frozenset(
    {
        RowStatus.DUPLICADO,
        RowStatus.OBRIGATORIO,
        RowStatus.CONFLITO,
        RowStatus.INVALIDA,
    }
)


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------


def validar_upload(arquivo):
    if arquivo is None:
        raise DomainError("Selecione um arquivo para importar.")
    nome = (getattr(arquivo, "name", "") or "").lower()
    if not nome.endswith(EXTENSOES_ACEITAS):
        raise DomainError("Formato não suportado. Envie um arquivo .xlsx ou .csv.")
    tamanho = getattr(arquivo, "size", 0) or 0
    if tamanho > TAMANHO_MAXIMO_BYTES:
        raise DomainError(
            "O arquivo excede o limite de {} MB.".format(
                TAMANHO_MAXIMO_BYTES // (1024 * 1024)
            )
        )
    if tamanho == 0:
        raise DomainError("O arquivo está vazio.")
    return nome


def ler_arquivo(arquivo):
    """
    Le o arquivo e devolve o lote cru, pronto para a sessao.

        {
          "aba": nome da aba lida (ou "" no CSV),
          "linha_cabecalho": numero da linha do cabecalho no arquivo,
          "colunas": [{"cabecalho", "campo", "rotulo"}...],
          "ignoradas": [{"cabecalho", "motivo"}...],
          "linhas": [{"numero", "sequencia", "nome", ...}, ...],
        }

    So tipos simples (texto e inteiro): o lote atravessa a sessao em JSON.
    """
    nome = validar_upload(arquivo)
    arquivo.seek(0)
    if nome.endswith(".csv"):
        return _montar_lote([("", _ler_csv(arquivo))])
    return _montar_lote(_ler_xlsx(arquivo))


def _ler_csv(arquivo):
    """
    CSV em UTF-8 (com ou sem BOM), separado por virgula ou ponto e virgula.

    Mesma decisao da importacao de alunos: nao ha recuo silencioso para
    cp1252. Um nome com acento lido errado seria gravado errado, e o
    administrador so perceberia na guia impressa.
    """
    bruto = arquivo.read()
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DomainError(
            "Não foi possível ler o arquivo: a codificação não é UTF-8. "
            "Salve novamente como 'CSV UTF-8' ou envie o .xlsx original."
        )
    if not texto.strip():
        raise DomainError("O arquivo está vazio.")
    amostra = texto[:8192]
    try:
        delimitador = csv.Sniffer().sniff(amostra, delimiters=",;").delimiter
    except csv.Error:
        delimitador = ";" if amostra.count(";") > amostra.count(",") else ","
    try:
        return list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    except csv.Error as exc:
        raise DomainError("Não foi possível interpretar o CSV: {}".format(exc))


# Mesma lista, e pelos mesmos motivos, de students.importers. Repetida aqui
# para que o dominio de ministros nao importe nada do dominio academico.
ERROS_DE_PLANILHA_INVALIDA = (
    zipfile.BadZipFile,
    KeyError,
    OSError,
    ValueError,
    ParseError,
    IndexError,
)


def _ler_xlsx(arquivo):
    """Devolve [(nome_da_aba, linhas)] de cada aba visivel."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    invalidos = ERROS_DE_PLANILHA_INVALIDA + (InvalidFileException,)
    planilha = None
    abas = []
    try:
        planilha = load_workbook(
            filename=io.BytesIO(arquivo.read()), read_only=True, data_only=True
        )
        for aba in planilha.worksheets:
            if getattr(aba, "sheet_state", "visible") != "visible":
                continue
            linhas = []
            for linha in aba.iter_rows(values_only=True):
                linhas.append(list(linha))
                if len(linhas) > MAX_LINHAS + LINHAS_PROCURANDO_CABECALHO:
                    break
            abas.append((aba.title, linhas))
    except invalidos as exc:
        logger.info("Planilha de ministros recusada (%s)", exc.__class__.__name__)
        raise DomainError(
            "Não foi possível ler a planilha. Confirme que o arquivo é um "
            ".xlsx válido e não está protegido por senha."
        )
    finally:
        if planilha is not None:
            planilha.close()
    if not abas:
        raise DomainError("A planilha não possui nenhuma aba visível.")
    return abas


def _texto(valor):
    """Celula como texto limpo. Inteiro do Excel nao vira "132921.0"."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return ""
    if isinstance(valor, datetime.datetime):
        return valor.date().isoformat()
    if isinstance(valor, datetime.date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _celula(linha, posicao):
    return linha[posicao] if posicao is not None and posicao < len(linha) else None


def _encontrar_cabecalho(linhas):
    """
    Posicao da linha de cabecalho e o mapa campo -> coluna.

    E a primeira linha, entre as do topo, que tem ao mesmo tempo uma coluna
    de nome e uma de inscricao CONAMAD.
    """
    for indice, linha in enumerate(linhas[:LINHAS_PROCURANDO_CABECALHO]):
        mapa = {}
        for posicao, valor in enumerate(linha):
            campo = ALIAS_COLUNAS.get(chave_de_texto(_texto(valor)))
            if campo and campo not in mapa:
                mapa[campo] = posicao
        if "nome" in mapa and "conamad" in mapa:
            return indice, mapa
    return None, None


def _montar_lote(abas):
    for nome_aba, linhas in abas:
        indice, mapa = _encontrar_cabecalho(linhas)
        if indice is not None:
            return _extrair(nome_aba, linhas, indice, mapa)
    raise DomainError(
        "Não encontrei o cabeçalho da relação. O arquivo precisa ter, nas "
        "primeiras {} linhas, as colunas \"Nome\" e \"Inscrição Conamad\".".format(
            LINHAS_PROCURANDO_CABECALHO
        )
    )


def _nascimento_em_tres_colunas(cabecalho, posicao):
    """
    Se "Data Nascimento" cobre tambem as duas colunas seguintes.

    Na relacao oficial o cabecalho esta mesclado de H a J, e a celula
    mesclada so tem valor na primeira coluna: as duas seguintes chegam
    vazias. E isso que caracteriza o formato dia | mes | ano.
    """
    return all(
        not _texto(_celula(cabecalho, posicao + deslocamento))
        for deslocamento in (1, 2)
    )


def _ler_nascimento(linha, posicao, tres_colunas):
    valor = _celula(linha, posicao)
    if isinstance(valor, (datetime.date, datetime.datetime)):
        return _texto(valor)
    texto = _texto(valor)
    if tres_colunas and texto.isdigit():
        mes = _texto(_celula(linha, posicao + 1))
        ano = _texto(_celula(linha, posicao + 2))
        if mes and ano:
            return "{}/{}/{}".format(texto, mes, ano)
    return texto


def _extrair(nome_aba, linhas, indice, mapa):
    cabecalho = linhas[indice]
    tres_colunas = "nascimento" in mapa and _nascimento_em_tres_colunas(
        cabecalho, mapa["nascimento"]
    )
    colunas_usadas = set(mapa.values())
    if tres_colunas:
        colunas_usadas |= {mapa["nascimento"] + 1, mapa["nascimento"] + 2}

    colunas = [
        {
            "cabecalho": _texto(cabecalho[posicao]),
            "campo": campo,
            "rotulo": ROTULO_DO_CAMPO[campo]
            + (" (dia, mês e ano em 3 colunas)" if campo == "nascimento" and tres_colunas else ""),
        }
        for campo, posicao in sorted(mapa.items(), key=lambda item: item[1])
    ]
    ignoradas = []
    for posicao, valor in enumerate(cabecalho):
        texto = _texto(valor)
        if posicao in colunas_usadas or not texto:
            continue
        chave = chave_de_texto(texto)
        motivo = COLUNAS_DERIVADAS.get(chave, "coluna não reconhecida")
        if chave in ALIAS_COLUNAS:
            motivo = "coluna repetida; vale a primeira"
        ignoradas.append({"cabecalho": texto, "motivo": motivo})

    registros = []
    for deslocamento, linha in enumerate(linhas[indice + 1 :], start=indice + 2):
        registro = {"numero": deslocamento}
        for campo, posicao in mapa.items():
            if campo == "nascimento":
                registro[campo] = _ler_nascimento(linha, posicao, tres_colunas)
            else:
                registro[campo] = _texto(_celula(linha, posicao))
        # Linha em branco e espaco da planilha, nao erro.
        if not any(registro.get(campo) for campo in mapa if campo != "sequencia"):
            continue
        registros.append(registro)

    if not registros:
        raise DomainError("O arquivo não possui nenhuma linha de dados.")
    if len(registros) > MAX_LINHAS:
        raise DomainError(
            "O arquivo possui {} linhas e o limite por importação é {}.".format(
                len(registros), MAX_LINHAS
            )
        )

    return {
        "aba": nome_aba,
        "linha_cabecalho": indice + 1,
        "colunas": colunas,
        "ignoradas": ignoradas,
        "linhas": registros,
    }


# ---------------------------------------------------------------------------
# Interpretacao de cada campo
# ---------------------------------------------------------------------------


def _interpretar_data(texto):
    texto = (texto or "").strip()
    if not texto:
        return None, None
    iso = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", texto)
    br = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", texto)
    try:
        if iso:
            return datetime.date(*(int(p) for p in iso.groups())), None
        if br:
            dia, mes, ano = (int(p) for p in br.groups())
            return datetime.date(ano, mes, dia), None
    except ValueError:
        pass
    return None, "Data de nascimento inválida ({}).".format(texto)


def _interpretar_ano(texto):
    texto = (texto or "").strip()
    if not texto:
        return None, None
    if re.fullmatch(r"\d{4}", texto):
        return int(texto), None
    return None, "Ano de consagração inválido ({}).".format(texto)


def _interpretar_titulo(texto):
    if not (texto or "").strip():
        return "", None
    titulo = TITULOS.get(chave_de_texto(texto))
    if titulo is None:
        return "", "Título não reconhecido ({}).".format(texto)
    return titulo, None


def _interpretar_sexo(texto):
    chave = chave_de_texto(texto)
    if not chave:
        return "", None
    if chave in SEXOS:
        return SEXOS[chave], None
    if chave in SEXOS_AMBIGUOS:
        return "", (
            "Sexo \"{}\" é ambíguo na relação (usado para masculino e para "
            "mulher); ficará em branco.".format(texto)
        )
    return "", "Sexo \"{}\" não reconhecido; ficará em branco.".format(texto)


# ---------------------------------------------------------------------------
# Analise
# ---------------------------------------------------------------------------


@dataclass
class LinhaAnalisada:
    numero: int
    sequencia: str
    nome: str
    conamad: str
    campo: str
    titulo: str
    status: str
    problemas: list = field(default_factory=list)
    avisos: list = field(default_factory=list)
    divergencias: list = field(default_factory=list)
    dados: dict = field(default_factory=dict)

    @property
    def importavel(self):
        return self.status == RowStatus.NOVO

    @property
    def existente(self):
        return self.status in STATUS_EXISTENTES

    @property
    def recusada(self):
        return self.status in STATUS_RECUSADOS

    @property
    def import_sequence(self):
        return int(self.sequencia) if self.sequencia.isdigit() else None


@dataclass
class ResultadoAnalise:
    linhas: list = field(default_factory=list)

    def _contar(self, *status):
        return sum(1 for linha in self.linhas if linha.status in status)

    @property
    def novos(self):
        return [linha for linha in self.linhas if linha.importavel]

    @property
    def tem_algo_a_importar(self):
        return bool(self.novos)

    def resumo(self):
        return {
            "total_linhas": len(self.linhas),
            "novos": self._contar(RowStatus.NOVO),
            "existentes": self._contar(RowStatus.EXISTENTE, RowStatus.DIVERGENTE),
            "divergentes": self._contar(RowStatus.DIVERGENTE),
            "duplicados": self._contar(RowStatus.DUPLICADO),
            "obrigatorios_ausentes": self._contar(RowStatus.OBRIGATORIO),
            "sem_identificador": sum(1 for linha in self.linhas if not linha.conamad),
            "conflitantes": self._contar(RowStatus.CONFLITO),
            "invalidas": self._contar(RowStatus.INVALIDA),
            "recusadas": self._contar(*STATUS_RECUSADOS),
            "com_aviso": sum(1 for linha in self.linhas if linha.avisos),
        }


def _divergencias(ministro, dados):
    """Campos em que o arquivo diz algo diferente do cadastro."""
    diferentes = []
    if " ".join(dados["full_name"].split()).casefold() != ministro.full_name.casefold():
        diferentes.append("nome")
    if dados["ministerial_title"] != ministro.ministerial_title:
        diferentes.append("título")
    # Campo vazio no arquivo e ausencia de informacao, nao divergencia.
    comparaveis = (
        ("field_registration", "matrícula do campo"),
        ("position", "cargo / função"),
        ("congregation", "congregação"),
        ("consecration_year", "ano de consagração"),
        ("birth_date", "data de nascimento"),
        ("sex", "sexo"),
    )
    for campo, rotulo in comparaveis:
        valor = dados.get(campo)
        if valor not in ("", None) and valor != getattr(ministro, campo):
            diferentes.append(rotulo)
    return diferentes


def analisar(linhas_cruas):
    """
    Classifica cada linha confrontando o arquivo com o banco.

    Somente leitura, em qualquer circunstancia. Tres consultas no total,
    independentemente do tamanho do arquivo.
    """
    preparadas = []
    for crua in linhas_cruas:
        preparadas.append(_preparar(crua))

    inscricoes = {p["dados"]["conamad_registration"] for p in preparadas} - {""}
    matriculas = {p["dados"]["field_registration"] for p in preparadas} - {""}

    existentes = {
        m.conamad_registration: m
        for m in Minister.objects.filter(conamad_registration__in=inscricoes)
    }
    donos_da_matricula = dict(
        Minister.objects.filter(field_registration__in=matriculas).values_list(
            "field_registration", "conamad_registration"
        )
    )
    nomes_cadastrados = {}
    for nome, inscricao in Minister.objects.values_list(
        "full_name", "conamad_registration"
    ):
        nomes_cadastrados.setdefault(chave_de_nome(nome), set()).add(inscricao)

    # Repeticoes dentro do proprio arquivo.
    ocorrencias_inscricao = {}
    inscricoes_por_matricula = {}
    inscricoes_por_nome = {}
    for p in preparadas:
        inscricao = p["dados"]["conamad_registration"]
        if inscricao:
            ocorrencias_inscricao[inscricao] = ocorrencias_inscricao.get(inscricao, 0) + 1
        matricula = p["dados"]["field_registration"]
        if matricula and inscricao:
            inscricoes_por_matricula.setdefault(matricula, set()).add(inscricao)
        nome = chave_de_nome(p["dados"]["full_name"])
        if nome and inscricao:
            inscricoes_por_nome.setdefault(nome, set()).add(inscricao)

    analisadas = []
    for p in preparadas:
        dados = p["dados"]
        inscricao = dados["conamad_registration"]
        matricula = dados["field_registration"]

        def montar(status, *, divergencias=None):
            return LinhaAnalisada(
                numero=p["numero"],
                sequencia=p["sequencia"],
                nome=dados["full_name"],
                conamad=inscricao,
                campo=matricula,
                titulo=p["titulo_original"],
                status=status,
                problemas=p["problemas"],
                avisos=p["avisos"],
                divergencias=divergencias or [],
                dados=dados,
            )

        if p["faltando"]:
            analisadas.append(montar(RowStatus.OBRIGATORIO))
            continue
        if p["problemas"]:
            analisadas.append(montar(RowStatus.INVALIDA))
            continue

        if ocorrencias_inscricao.get(inscricao, 0) > 1:
            p["problemas"].append(
                "A inscrição {} aparece {} vezes no arquivo.".format(
                    inscricao, ocorrencias_inscricao[inscricao]
                )
            )
            analisadas.append(montar(RowStatus.DUPLICADO))
            continue

        if matricula:
            dono = donos_da_matricula.get(matricula)
            if dono and dono != inscricao:
                p["problemas"].append(
                    "A matrícula do campo {} já pertence a outro ministro "
                    "(inscrição {}).".format(matricula, dono)
                )
                analisadas.append(montar(RowStatus.CONFLITO))
                continue
            if len(inscricoes_por_matricula.get(matricula, ())) > 1:
                p["problemas"].append(
                    "A matrícula do campo {} aparece no arquivo para mais de "
                    "uma inscrição.".format(matricula)
                )
                analisadas.append(montar(RowStatus.CONFLITO))
                continue

        ministro = existentes.get(inscricao)
        if ministro is not None:
            diferentes = _divergencias(ministro, dados)
            if diferentes:
                analisadas.append(
                    montar(RowStatus.DIVERGENTE, divergencias=diferentes)
                )
            else:
                analisadas.append(montar(RowStatus.EXISTENTE))
            continue

        # Homonimo nao impede: duas pessoas podem ter o mesmo nome. Mas pode
        # ser a mesma pessoa com a inscricao digitada errado, e isso o
        # administrador precisa ver antes de confirmar.
        nome = chave_de_nome(dados["full_name"])
        outras = (nomes_cadastrados.get(nome, set()) | inscricoes_por_nome.get(nome, set())) - {inscricao}
        if outras:
            p["avisos"].append(
                "Mesmo nome de outro registro (inscrição {}).".format(
                    ", ".join(sorted(outras))
                )
            )
        analisadas.append(montar(RowStatus.NOVO))

    return ResultadoAnalise(linhas=analisadas)


def _preparar(crua):
    """Interpreta uma linha crua, sem consultar o banco."""
    problemas = []
    avisos = []

    nome = normalizar_nome(crua.get("nome"))
    inscricao = normalizar_registro(crua.get("conamad"))
    matricula = normalizar_registro(crua.get("campo"))
    titulo_original = (crua.get("titulo") or "").strip()

    faltando = []
    if not nome:
        faltando.append("nome")
    if not inscricao:
        faltando.append("inscrição CONAMAD")
    if not titulo_original:
        faltando.append("título")
    if faltando:
        problemas.append("Sem {}.".format(", ".join(faltando)))

    if inscricao and not registro_valido(inscricao):
        problemas.append(
            "Inscrição CONAMAD com formato inválido ({}).".format(inscricao[:40])
        )
    if matricula and not registro_valido(matricula):
        problemas.append(
            "Matrícula do campo com formato inválido ({}).".format(matricula[:40])
        )
    if len(nome) > 200:
        problemas.append("Nome com mais de 200 caracteres.")

    titulo, erro = _interpretar_titulo(titulo_original)
    if erro:
        problemas.append(erro)

    ano, erro = _interpretar_ano(crua.get("consagracao"))
    if erro:
        problemas.append(erro)
    elif ano is not None and not (services.ANO_MINIMO <= ano <= services._ano_atual()):
        problemas.append("Ano de consagração fora do intervalo ({}).".format(ano))
        ano = None

    nascimento, erro = _interpretar_data(crua.get("nascimento"))
    if erro:
        problemas.append(erro)
    elif nascimento is not None and not (
        nascimento.year >= services.ANO_MINIMO
        and nascimento <= timezone.localdate()
    ):
        problemas.append("Data de nascimento fora do intervalo.")
        nascimento = None

    sexo, aviso = _interpretar_sexo(crua.get("sexo"))
    if aviso:
        avisos.append(aviso)

    cargo = normalizar_nome(crua.get("cargo"))[:100]
    congregacao = normalizar_nome(crua.get("congregacao"))[:120]

    return {
        "numero": crua.get("numero", 0),
        "sequencia": (crua.get("sequencia") or "").strip(),
        "titulo_original": titulo_original,
        "faltando": faltando,
        "problemas": problemas,
        "avisos": avisos,
        "dados": {
            "full_name": nome,
            "conamad_registration": inscricao[:TAMANHO_MAXIMO_REGISTRO],
            "field_registration": matricula[:TAMANHO_MAXIMO_REGISTRO],
            "ministerial_title": titulo,
            "position": cargo,
            "congregation": congregacao,
            "consecration_year": ano,
            "birth_date": nascimento,
            "sex": sexo,
            "notes": (crua.get("observacoes") or "").strip(),
        },
    }


# ---------------------------------------------------------------------------
# Confirmacao
# ---------------------------------------------------------------------------


def confirmar(linhas_cruas, *, criar_acesso, actor=None, request=None):
    """
    Grava os ministros NOVOS do lote.

    A analise e refeita aqui: o que o navegador mostrou nao decide nada, e o
    formulario de confirmacao so carrega o token do lote e a escolha de criar
    acesso. Tudo numa transacao: ou o lote inteiro entra, ou nada entra.

    Com criar_acesso, cada ministro novo ganha acesso ao portal com a senha
    padrao (DEFAULT_MINISTER_PASSWORD) e a troca obrigatoria no primeiro
    login. A ausencia da variavel recusa a operacao ANTES de qualquer escrita.
    """
    hash_padrao = services.hash_da_senha_padrao() if criar_acesso else None

    try:
        with transaction.atomic():
            analise = analisar(linhas_cruas)
            if not analise.tem_algo_a_importar:
                raise DomainError("Não há nenhum ministro novo para importar.")

            criados = 0
            acessos = 0
            for linha in analise.novos:
                ministro = services.create_minister(
                    dados=linha.dados,
                    source=MinisterSource.IMPORT,
                    import_sequence=linha.import_sequence,
                    actor=actor,
                    request=request,
                )
                criados += 1
                if hash_padrao is not None:
                    services._criar_acesso(
                        ministro,
                        hash_pronto=hash_padrao,
                        must_change_password=True,
                        actor=actor,
                        request=request,
                    )
                    acessos += 1

            resumo = analise.resumo()
            record(
                AuditEvent.MINISTER_IMPORT_COMPLETED,
                request=request,
                actor=actor,
                entity_type="Minister",
                metadata={
                    # Apenas contagens. Os eventos individuais ja dao a
                    # rastreabilidade por ministro.
                    "rows_received": resumo["total_linhas"],
                    "ministers_created": criados,
                    "accounts_created": acessos,
                    "already_existing": resumo["existentes"],
                    "rejected": resumo["recusadas"],
                },
            )
    except IntegrityError:
        # Outro administrador gravou o mesmo ministro entre a analise e a
        # escrita. A transacao inteira voltou; o preview novo vai mostrar.
        logger.info("Importacao de ministros desfeita por conflito de unicidade")
        raise DomainError(
            "Outro cadastro foi gravado ao mesmo tempo e a importação foi "
            "desfeita por inteiro. Revise o preview e confirme novamente."
        )

    return {
        "criados": criados,
        "acessos": acessos,
        "existentes": resumo["existentes"],
        "recusadas": resumo["recusadas"],
        "total_linhas": resumo["total_linhas"],
    }
