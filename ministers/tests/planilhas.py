"""
Arquivos de teste no formato REAL da relacao de ministros.

A planilha oficial ("Relatorio Ministros 2026.xlsx") nao entra no
repositorio: tem dado pessoal de 84 pessoas. Esta funcao reproduz a
ESTRUTURA dela, celula por celula, com nomes e numeros ficticios:

    aba "Ministros Campo Sumaré"
    C2:L4 mesclado com o titulo "Ministro Ad Brás Sumaré"
    linha 6: B "nº", C "Nome", E "Inscrição Conamad", F "Titulo",
             G "Consagração", H "Data Nascimento" (mesclado H6:J6),
             K "IDADE ", L "Sexo"
    linha 7 em branco
    dados a partir da linha 8: nº sequencial, nome em maiusculas, inscricao
    como numero, titulo abreviado, ano de consagracao, dia | mes | ano em
    tres colunas, IDADE como formula (=2026-J8) e sexo H/M/F
"""

import io

from django.core.files.uploadedfile import SimpleUploadedFile

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# (nome, inscricao, titulo, consagracao, (dia, mes, ano), sexo)
LINHAS_EXEMPLO = [
    ("ABEL FICTÍCIO DOS SANTOS", 132901, "EV", 2025, (9, 4, 2000), "M"),
    ("BENTO EXEMPLO DA SILVA", 129102, "PR", 2016, (21, 9, 1975), "H"),
    ("CÁSSIA MODELO PEREIRA", 132903, "EV", 2024, (1, 7, 1980), "F"),
    ("DÉBORA TESTE ALVES", 15504, "PRª", "2009", (5, 3, 1961), "M"),
    ("ELIAS AMOSTRA COSTA", 129105, "Ev", 2024, (1, 12, 1982), "H"),
    ("FÁTIMA PROTÓTIPO LIMA", 132906, "Mis", 2023, (30, 1, 1970), "M"),
]


def planilha_relacao(linhas=None, nome="Relatorio Ministros 2026.xlsx", aba="Ministros Campo Sumaré"):
    from openpyxl import Workbook

    linhas = LINHAS_EXEMPLO if linhas is None else linhas
    livro = Workbook()
    folha = livro.active
    folha.title = aba
    folha["C2"] = "           Ministro Ad Brás Sumaré"
    folha.merge_cells("C2:L4")
    cabecalho = {
        "B6": "nº",
        "C6": "Nome",
        "E6": "Inscrição Conamad",
        "F6": "Titulo",
        "G6": "Consagração",
        "H6": "Data Nascimento",
        "K6": "IDADE ",
        "L6": "Sexo",
    }
    for celula, valor in cabecalho.items():
        folha[celula] = valor
    folha.merge_cells("H6:J6")

    for indice, linha in enumerate(linhas, start=1):
        n = 7 + indice
        nome_ministro, inscricao, titulo, consagracao, nascimento, sexo = linha
        folha["B{}".format(n)] = indice
        folha["C{}".format(n)] = nome_ministro
        folha["E{}".format(n)] = inscricao
        folha["F{}".format(n)] = titulo
        folha["G{}".format(n)] = consagracao
        if nascimento:
            dia, mes, ano = nascimento
            folha["H{}".format(n)] = dia
            folha["I{}".format(n)] = mes
            folha["J{}".format(n)] = ano
            folha["K{}".format(n)] = "=2026-J{}".format(n)
        folha["L{}".format(n)] = sexo

    buffer = io.BytesIO()
    livro.save(buffer)
    return SimpleUploadedFile(nome, buffer.getvalue(), content_type=TIPO_XLSX)


def csv_relacao(texto, nome="relacao.csv", bom=True):
    dados = texto.encode("utf-8")
    if bom:
        dados = b"\xef\xbb\xbf" + dados
    return SimpleUploadedFile(nome, dados, content_type="text/csv")


# A mesma relacao exportada pelo Excel como "CSV UTF-8": ponto e virgula,
# as linhas de titulo e a linha em branco acima dos dados preservadas, e a
# data de nascimento ainda em tres colunas.
CSV_EXPORTADO_DO_EXCEL = (
    ";;           Ministro Ad Brás Sumaré;;;;;;;;;\n"
    ";;;;;;;;;;;\n"
    ";;;;;;;;;;;\n"
    ";;;;;;;;;;;\n"
    ";nº;Nome;;Inscrição Conamad;Titulo;Consagração;Data Nascimento;;;IDADE ;Sexo\n"
    ";;;;;;;;;;;\n"
    ";1;ABEL FICTÍCIO DOS SANTOS;;132901;EV;2025;9;4;2000;26;M\n"
    ";2;BENTO EXEMPLO DA SILVA;;129102;PR;2016;21;9;1975;51;H\n"
    ";3;CÁSSIA MODELO PEREIRA;;132903;EV;2024;1;7;1980;46;F\n"
)
